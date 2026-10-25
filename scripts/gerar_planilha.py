"""Gera uma planilha enxuta (XLSX) com o essencial da base.

Abas: Resumo (fórmulas), Acórdãos (sem o voto), Desfechos oficiais, Estoque, Revisão LGPD.
Uso: python scripts/gerar_planilha.py DATA_DIR SAIDA.xlsx [REVISAO.csv]
"""

import csv
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from carf.qa import consulta
from carf.qa import gerenciais as qg
from carf.store import Store

FONTE = "Arial"
CAB = PatternFill("solid", start_color="1F3864")
ENTRADA = PatternFill("solid", start_color="FFF2CC")
LINK_BASE = "https://acordaos.economia.gov.br/acordaos2/pdfs/processados/"
fino = Side(style="thin", color="BFBFBF")


def novo():
    wb = Workbook()
    wb._named_styles["Normal"].font = Font(name=FONTE, size=10)
    return wb


def cabecalho(ws, colunas, larguras, linha=1):
    for i, (c, w) in enumerate(zip(colunas, larguras, strict=True), 1):
        cel = ws.cell(row=linha, column=i, value=c)
        cel.font = Font(name=FONTE, bold=True, color="FFFFFF", size=10)
        cel.fill = CAB
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = f"A{linha + 1}"  # texto: não cria célula (criaria linha em branco)
    ws.auto_filter.ref = f"A{linha}:{get_column_letter(len(colunas))}{linha}"


def dados_acordaos(con, raiz: Path) -> pd.DataFrame:
    df = con.execute(
        """SELECT id, numero_decisao, numero_processo, turma, secao, relator, dt_sessao, ano,
                  substr(ementa, 1, 200) AS ementa, decisao, nome_arquivo_pdf FROM acordaos
           ORDER BY ano NULLS LAST, dt_sessao, numero_decisao"""
    ).df()
    df["resultado_heuristico"] = df.decisao.map(consulta.classificar_resultado)
    df = df.drop(columns="decisao")
    df["link"] = LINK_BASE + df.nome_arquivo_pdf.fillna("")
    df.loc[df.nome_arquivo_pdf.isna(), "link"] = None
    df = df.drop(columns="nome_arquivo_pdf")
    j = pd.read_parquet(raiz / "gerenciais/julgamentos/julgamentos_2025.parquet",
                        columns=["numero_decisao", "tributo", "desfecho", "tipo_resultado"])
    j = j[j.tipo_resultado == "Acórdão"]
    grp = j.groupby("numero_decisao").agg(tributo=("tributo", "first"), n=("desfecho", "size"),
                                          desfecho=("desfecho", "first"))
    grp.loc[grp.n > 1, "desfecho"] = "VÁRIOS RECURSOS"
    df = df.merge(grp[["tributo", "desfecho"]].rename(columns={"desfecho": "resultado_oficial"}),
                  left_on="numero_decisao", right_index=True, how="left")
    return df


def aba_acordaos(wb, df):
    ws = wb.create_sheet("Acórdãos")
    cols = ["ID", "Nº decisão", "Processo", "Turma", "Seção", "Relator", "Data da sessão", "Ano",
            "Ementa (200 car.)", "Resultado (heurística)", "Link do PDF", "Tributo (oficial 2025)",
            "Resultado oficial (2025)"]
    cabecalho(ws, cols, [11, 14, 22, 38, 26, 28, 13, 7, 60, 18, 40, 22, 22])
    ordem = ["id", "numero_decisao", "numero_processo", "turma", "secao", "relator", "dt_sessao", "ano",
             "ementa", "resultado_heuristico", "link", "tributo", "resultado_oficial"]
    for linha in df[ordem].itertuples(index=False):
        ws.append([None if pd.isna(v) else (v.item() if hasattr(v, "item") else v) for v in linha])
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=7).number_format = "dd/mm/yyyy"
    return ws


def aba_desfechos(wb, raiz: Path):
    con = qg.conectar(raiz)
    t = qg.desfecho_por(con, "tributo")
    ws = wb.create_sheet("Desfechos 2025")
    ws["A1"] = "Resultado OFICIAL dos recursos julgados em 2025 (acórdãos), por tributo"
    ws["A1"].font = Font(name=FONTE, bold=True, size=12)
    ws["A2"] = ("Fonte: dado aberto CARF - Julgamentos de processos (carf_julgamentos_2025.zip). "
                "Contagens são entradas (azul); totais e percentuais são fórmulas.")
    ws["A2"].font = Font(name=FONTE, italic=True, size=9)
    cols = [c for c in t.columns if c not in ("tributo", "total")]
    cabecalho(ws, ["Tributo", *cols, "Total", "% provido (total + parcial)", "% negado"],
              [34] + [14] * len(cols) + [10, 16, 10], linha=4)
    for i, r in enumerate(t.itertuples(index=False), 5):
        ws.cell(row=i, column=1, value=r.tributo)
        for j, c in enumerate(cols, 2):
            cel = ws.cell(row=i, column=j, value=int(getattr(r, c)))
            cel.font = Font(name=FONTE, color="0000FF")
            cel.number_format = "#,##0"
    ult = 4 + len(t)
    ini, fim = get_column_letter(2), get_column_letter(1 + len(cols))
    col_tot = 2 + len(cols)
    idx = {c: get_column_letter(2 + k) for k, c in enumerate(cols)}
    for i in range(5, ult + 1):
        ws.cell(row=i, column=col_tot, value=f"=SUM({ini}{i}:{fim}{i})").number_format = "#,##0"
        pv = f"({idx['PROVIDO']}{i}+{idx['PROVIDO_PARCIAL']}{i})"
        ws.cell(row=i, column=col_tot + 1, value=f"=IFERROR({pv}/{get_column_letter(col_tot)}{i},0)").number_format = "0.0%"
        ws.cell(row=i, column=col_tot + 2,
                value=f"=IFERROR({idx['NEGADO']}{i}/{get_column_letter(col_tot)}{i},0)").number_format = "0.0%"
    tot = ult + 1
    ws.cell(row=tot, column=1, value="Total").font = Font(name=FONTE, bold=True)
    for j in range(2, col_tot + 1):
        L = get_column_letter(j)
        c = ws.cell(row=tot, column=j, value=f"=SUM({L}5:{L}{ult})")
        c.font = Font(name=FONTE, bold=True)
        c.number_format = "#,##0"
        c.border = Border(top=fino)
    L = get_column_letter(col_tot)
    pv = f"({idx['PROVIDO']}{tot}+{idx['PROVIDO_PARCIAL']}{tot})"
    ws.cell(row=tot, column=col_tot + 1, value=f"=IFERROR({pv}/{L}{tot},0)").number_format = "0.0%"
    ws.cell(row=tot, column=col_tot + 2, value=f"=IFERROR({idx['NEGADO']}{tot}/{L}{tot},0)").number_format = "0.0%"
    ws.freeze_panes = "B5"
    return ws, tot, get_column_letter(col_tot)


def aba_estoque(wb, raiz: Path):
    con = qg.conectar(raiz)
    ws = wb.create_sheet("Estoque")
    ref = qg.estoque_por(con, "tributo")
    ano, mes = map(int, ref.attrs["referencia"].split("-"))
    ws["A1"] = f"Processos em estoque no CARF, referência {mes:02d}/{ano}"
    ws["A1"].font = Font(name=FONTE, bold=True, size=12)
    ws["A2"] = ("Fonte: dado aberto CARF - Estoque de processos (carf_estoque_2026_2semestreparcial.zip). "
                "Contagem de processos distintos (azul); totais são fórmulas.")
    ws["A2"].font = Font(name=FONTE, italic=True, size=9)
    blocos = [("Por tributo", ref, "tributo"),
              ("Por fase processual", qg.estoque_por(con, "fase_processual"), "fase_processual"),
              ("Por tipo de contribuinte", qg.estoque_por(con, "tipo_contribuinte"), "tipo_contribuinte")]
    col = 1
    for titulo, df, nome in blocos:
        ws.cell(row=4, column=col, value=titulo).font = Font(name=FONTE, bold=True)
        for k, h in enumerate([nome.replace("_", " ").capitalize(), "Processos", "% do total"]):
            cel = ws.cell(row=5, column=col + k, value=h)
            cel.font = Font(name=FONTE, bold=True, color="FFFFFF")
            cel.fill = CAB
        for i, r in enumerate(df.itertuples(index=False), 6):
            ws.cell(row=i, column=col, value=r[0])
            c = ws.cell(row=i, column=col + 1, value=int(r[1]))
            c.font = Font(name=FONTE, color="0000FF")
            c.number_format = "#,##0"
        fim = 5 + len(df)
        Lp = get_column_letter(col + 1)
        for i in range(6, fim + 1):
            ws.cell(row=i, column=col + 2, value=f"=IFERROR({Lp}{i}/{Lp}${fim + 1},0)").number_format = "0.0%"
        ws.cell(row=fim + 1, column=col, value="Total").font = Font(name=FONTE, bold=True)
        t = ws.cell(row=fim + 1, column=col + 1, value=f"=SUM({Lp}6:{Lp}{fim})")
        t.font = Font(name=FONTE, bold=True)
        t.number_format = "#,##0"
        ws.column_dimensions[get_column_letter(col)].width = 36
        ws.column_dimensions[Lp].width = 12
        ws.column_dimensions[get_column_letter(col + 2)].width = 11
        col += 4
    # série mensal do semestre carregado
    serie = con.execute("SELECT ano_referencia, mes_referencia, count(DISTINCT numero_processo) "
                        "FROM estoque WHERE ano_referencia = ? GROUP BY 1, 2 ORDER BY 2", [ano]).fetchall()
    base = 6 + max(len(b[1]) for b in blocos) + 3
    ws.cell(row=base, column=1, value="Estoque total por mês").font = Font(name=FONTE, bold=True)
    for k, h in enumerate(["Mês", "Processos"]):
        cel = ws.cell(row=base + 1, column=1 + k, value=h)
        cel.font = Font(name=FONTE, bold=True, color="FFFFFF")
        cel.fill = CAB
    for i, (a, m, n) in enumerate(serie, base + 2):
        ws.cell(row=i, column=1, value=f"{int(m):02d}/{int(a)}")
        c = ws.cell(row=i, column=2, value=int(n))
        c.font = Font(name=FONTE, color="0000FF")
        c.number_format = "#,##0"
    return ws


def aba_revisao(wb, caminho: Path):
    ws = wb.create_sheet("Revisão LGPD")
    ws["A1"] = "Revisão humana da anonimização: 100 acórdãos (semente 1)"
    ws["A1"].font = Font(name=FONTE, bold=True, size=12)
    ws["A2"] = ("Leia o texto no arquivo revisao/amostra-*.md e preencha as células amarelas com a CONTAGEM. "
                "Ex.: 0 | 0 | 1. Nomes e identificadores que ainda aparecem = vazamento; "
                "máscara indevida = termo jurídico ou empresa mascarado sem necessidade.")
    ws["A2"].font = Font(name=FONTE, italic=True, size=9)
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:H2")
    ws.row_dimensions[2].height = 38
    cols = ["ID", "Citação", "Palavras", "Máscaras", "Nomes remanescentes", "Identificadores remanescentes",
            "Máscaras indevidas", "Observações"]
    cabecalho(ws, cols, [11, 70, 10, 10, 14, 16, 12, 40], linha=4)
    with caminho.open(encoding="utf-8") as f:
        linhas = list(csv.DictReader(f))
    for i, r in enumerate(linhas, 5):
        ws.cell(row=i, column=1, value=r["id"])
        ws.cell(row=i, column=2, value=r["citacao"])
        ws.cell(row=i, column=3, value=int(r["palavras"])).number_format = "#,##0"
        ws.cell(row=i, column=4, value=int(r["total_mascaras"]))
        for c in range(5, 9):
            ws.cell(row=i, column=c).fill = ENTRADA
    return ws, 4 + len(linhas)


def aba_resumo(wb, n_acord, ult_acord, tot_des, col_des, ult_rev):
    ws = wb["Resumo"]
    ws["A1"] = "CARF Intelligence: resumo da base"
    ws["A1"].font = Font(name=FONTE, bold=True, size=14)
    ws["A2"] = "Valores em azul são entradas medidas (fonte na coluna C). Em preto, fórmulas sobre as outras abas."
    ws["A2"].font = Font(name=FONTE, italic=True, size=9)
    ac = f"'Acórdãos'!$H$2:$H${ult_acord}"
    linhas = [
        ("BASE", None, None),
        ("Acórdãos e resoluções na base (2023 em diante)", f"=COUNTA('Acórdãos'!$A$2:$A${ult_acord})", "aba Acórdãos"),
        ("  sessões em 2023", f"=COUNTIF({ac},2023)", "aba Acórdãos"),
        ("  sessões em 2024", f"=COUNTIF({ac},2024)", "aba Acórdãos"),
        ("  sessões em 2025", f"=COUNTIF({ac},2025)", "aba Acórdãos"),
        ("  sessões em 2026 (parcial)", f"=COUNTIF({ac},2026)", "aba Acórdãos"),
        ("  ano de sessão inválido ou ausente", f"=COUNTBLANK({ac})", "aba Acórdãos"),
        ("Recursos julgados em 2025 (oficial)", f"='Desfechos 2025'!{col_des}{tot_des}", "aba Desfechos 2025"),
        ("Processos em estoque (mês mais recente)", "=Estoque!B" + str(0), "aba Estoque"),
        ("FONTES E ATUALIZAÇÃO", None, None),
        ("Índice de acórdãos: última indexação", "24/10/2026", "campo dt_index_tdt do índice Solr do CARF, medido em 25/10/2026"),
        ("Julgamentos (dados abertos): última alteração em arquivo", "02/04/2026", "página do conjunto no dados.gov.br"),
        ("Estoque (dados abertos): última alteração em arquivo", "05/10/2026", "página do conjunto no dados.gov.br"),
        ("PDF gerencial: atualização da apresentação", "02/10/2026", "dados-abertos-gerenciais202609_v2.pdf"),
        ("QUALIDADE MEDIDA", None, None),
        ("Resultado (heurística): acurácia geral vs. oficial", "92,7% a 97,6%", "3 amostras de 2025, ~1,7 mil decisões cada, scripts/avaliar_resultado.py"),
        ("Resultado (heurística): decisões sem classificação (OUTRO)", "0,9% a 4,6%", "idem"),
        ("Máscaras de nome por mil palavras (ementa+decisão)", "0,18", "300 acórdãos, medição de 25/10/2026"),
        ("Taxa de vazamento real", "pendente", "depende da revisão humana (aba Revisão LGPD)"),
        ("Documentos já revisados (aba Revisão LGPD)", f"=COUNT('Revisão LGPD'!$E$5:$E${ult_rev})", "preenchimento manual"),
        ("  com nome ou identificador remanescente", f"=SUMPRODUCT(--(('Revisão LGPD'!$E$5:$E${ult_rev}+'Revisão LGPD'!$F$5:$F${ult_rev})>0))", "preenchimento manual"),
        ("  taxa de documentos com vazamento", None, "preenchimento manual"),
        ("LIMITES", None, None),
        ("Coluna 'Resultado (heurística)'", "indicativa", "é heurística sobre o texto da decisão; o oficial só existe para 2025 (outros anos: arquivos em carf.economia.gov.br)"),
        ("Anonimização", "reduz risco", "não elimina; ver docs/lgpd/ e a aba Revisão LGPD"),
        ("Esta planilha", "não é parecer", "apoio à pesquisa; confira sempre o acórdão de origem (coluna Link do PDF)"),
    ]
    cabecalho(ws, ["Item", "Valor", "Fonte / observação"], [58, 18, 90], linha=4)
    r = 5
    posicoes = {}
    for item, valor, fonte in linhas:
        if valor is None and fonte is None:
            cel = ws.cell(row=r, column=1, value=item)
            cel.font = Font(name=FONTE, bold=True)
            for c in range(1, 4):
                ws.cell(row=r, column=c).fill = PatternFill("solid", start_color="D9E1F2")
        else:
            ws.cell(row=r, column=1, value=item)
            if valor is not None:
                v = ws.cell(row=r, column=2, value=valor)
                v.alignment = Alignment(horizontal="right")
                if isinstance(valor, str) and not valor.startswith("="):
                    v.font = Font(name=FONTE, color="0000FF")
                else:
                    v.number_format = "#,##0"
            ws.cell(row=r, column=3, value=fonte)
            posicoes[item] = r
        r += 1
    # estoque: soma do bloco "por tributo"
    return ws, posicoes


def main(data_dir: str, saida: str, revisao: str | None):
    raiz = Path(data_dir)
    con = Store(raiz).connect()
    wb = novo()
    wb.active.title = "Resumo"
    df = dados_acordaos(con, raiz)
    ws_a = aba_acordaos(wb, df)
    _, tot_des, col_des = aba_desfechos(wb, raiz)
    ws_e = aba_estoque(wb, raiz)
    ult_rev = 4
    if revisao:
        _, ult_rev = aba_revisao(wb, Path(revisao))
    ws, pos = aba_resumo(wb, len(df), ws_a.max_row, tot_des, col_des, ult_rev)
    # estoque total = linha "Total" do bloco por tributo
    tot_est = next(r for r in range(6, ws_e.max_row + 1) if ws_e.cell(row=r, column=1).value == "Total")
    ws.cell(row=pos["Processos em estoque (mês mais recente)"], column=2, value=f"=Estoque!B{tot_est}")
    lin = pos["  taxa de documentos com vazamento"]
    ref_n = pos["Documentos já revisados (aba Revisão LGPD)"]
    ref_k = pos["  com nome ou identificador remanescente"]
    c = ws.cell(row=lin, column=2, value=f'=IFERROR(B{ref_k}/B{ref_n},"-")')
    c.number_format = "0.0%"
    c.alignment = Alignment(horizontal="right")
    for linha in ws.iter_rows(min_row=5, max_row=ws.max_row, min_col=3, max_col=3):
        for cel in linha:
            cel.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(saida)
    print(f"{saida}: {len(df)} acórdãos")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2], sys.argv[3] if len(sys.argv) > 3 else None)
